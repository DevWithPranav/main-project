"""Event exports (Build Plan M5, decision #3): CSV, GeoJSON, Excel (openpyxl), PDF (ReportLab).

All take the API event dicts (events.to_api) and return bytes. Positions are map metres (SRID 0),
so the GeoJSON carries a `crs_note` instead of pretending to be WGS84.
"""

import csv
import io
import json
from collections import Counter
from datetime import datetime, timezone

PDF_CHUNK = 40  # event rows per PDF table (about one landscape A4 page at 7 pt)
COLUMNS = ["event_id", "session_id", "kind", "type", "condition", "status", "review", "cls", "track_ids",
           "occurred_at", "t_s", "x", "y", "lane_id", "zone_id", "confidence", "tags", "value", "clip_url"]


def flat(e: dict) -> dict:
    return {
        "event_id": e["event_id"], "session_id": e.get("session_id"), "kind": e.get("kind", "violation"),
        "type": e["type"], "condition": e.get("condition"), "status": e["status"],
        "review": (e.get("review") or {}).get("outcome", ""), "cls": e.get("cls", ""),
        "track_ids": " ".join(map(str, e.get("track_ids", []))), "occurred_at": e.get("occurred_at"),
        "t_s": e.get("flag_s", e.get("t_s")), "x": e["x"], "y": e["y"], "lane_id": e.get("lane_id"),
        "zone_id": e.get("zone_id"), "confidence": e["confidence"], "tags": " ".join(e.get("tags", [])),
        "value": json.dumps(e.get("value", {k: e[k] for k in ("severity_score", "severity_band") if k in e})),
        "clip_url": (e.get("evidence") or {}).get("clip_url", ""),
    }


def to_csv(events: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS)
    w.writeheader()
    for e in events:
        w.writerow(flat(e))
    return buf.getvalue().encode("utf-8-sig")


def to_geojson(events: list[dict]) -> bytes:
    fc = {"type": "FeatureCollection",
          "crs_note": "coordinates are map metres (CARLA world / site frame, SRID 0), not WGS84",
          "features": [{"type": "Feature", "id": e["event_id"],
                        "geometry": {"type": "Point", "coordinates": [e["x"], e["y"]]},
                        "properties": {k: v for k, v in e.items() if k not in ("x", "y")}} for e in events]}
    return json.dumps(fc).encode()


def summary(events: list[dict]) -> list[tuple[str, list[tuple[str, int]]]]:
    return [("Type", sorted(Counter(e["type"] for e in events).items())),
            ("Condition", sorted(Counter(e.get("condition") or "-" for e in events).items())),
            ("Status", sorted(Counter(e["status"] for e in events).items())),
            ("Review", sorted(Counter((e.get("review") or {}).get("outcome", "none") for e in events).items()))]


def to_xlsx(events: list[dict], filters: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Traffic violation report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    ws.append(["Filters", json.dumps({k: v for k, v in filters.items() if v})])
    ws.append(["Events", len(events)])
    for title, rows in summary(events):
        ws.append([])
        ws.append([f"By {title.lower()}", "Count"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for k, n in rows:
            ws.append([k, n])
    ws.column_dimensions["A"].width = 28
    ev = wb.create_sheet("Events")
    ev.append(COLUMNS)
    for c in ev[1]:
        c.font = Font(bold=True)
    for e in events:
        r = flat(e)
        ev.append([r[k] for k in COLUMNS])
    ev.freeze_panes = "A2"
    ev.auto_filter.ref = ev.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(events: list[dict], filters: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="Traffic violation report")
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 7),
                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dde3ea")),
                       ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP")])
    story = [Paragraph("Traffic violation report", styles["Title"]),
             Paragraph(f"Generated {datetime.now(timezone.utc).isoformat(timespec='seconds')} &middot; "
                       f"{len(events)} events &middot; filters: "
                       f"{json.dumps({k: v for k, v in filters.items() if v}) or 'none'}", styles["Normal"]),
             Spacer(1, 4 * mm)]
    blocks = []
    for title, rows in summary(events):
        t = Table([[title, "Count"]] + [[k, n] for k, n in rows] if rows else [[title, "Count"]])
        t.setStyle(grid)
        blocks.append(t)
    row = Table([blocks], hAlign="LEFT")
    row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("RIGHTPADDING", (0, 0), (-1, -1), 8)]))
    story += [row, Spacer(1, 5 * mm)]
    cols = ["event_id", "type", "condition", "status", "review", "cls", "track_ids", "occurred_at", "x", "y",
            "lane_id", "confidence"]
    data = [cols]
    for e in events:
        r = flat(e)
        r["occurred_at"] = (r["occurred_at"] or "")[:19].replace("T", " ")
        data.append([str(r[k]) if r[k] is not None else "" for k in cols])
    # One table per PDF_CHUNK rows: ReportLab's split of one long table grows faster than linear
    # (10 000 rows took 10.97 s as one table on the dev laptop, 2026-10-10).
    widths = None
    for i in range(1, max(len(data), 2), PDF_CHUNK):
        t = Table([cols] + data[i:i + PDF_CHUNK], repeatRows=1, colWidths=widths)
        t.setStyle(grid)
        story.append(t)
    doc.build(story)
    return buf.getvalue()
