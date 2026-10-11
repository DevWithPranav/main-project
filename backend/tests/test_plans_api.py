"""Traffic plans API (routers/plans.py, ml/planning/plans.py): catalogue, locate, preview, save + versions, analysis."""

import os
import time

FLIGHT = os.getenv("TEST_FLIGHT", "20261009_201727")
SCENE = "Town05"


def doc(**kw):
    d = {"name": "api test plan", "scene": SCENE, "description": "", "session_id": None, "objectives": [], "props": [], "rules": [], "scenarios": []}
    d.update(kw)
    return d


def lane_xy(client, auth, x=-84.0, y=-144.0):
    return client.get("/api/plans/locate", headers=auth("planner"), params={"scene": SCENE, "x": x, "y": y}).json()


def test_catalogue_and_locate(client, auth):
    c = client.get("/api/plans/catalogue", headers=auth("officer")).json()
    assert {"traffic_signal", "speed_limit", "road_closure"} <= set(c["kinds"]) and c["templates"] and "preliminary" in c["evidence_levels"]
    r = lane_xy(client, auth)
    assert r["on_road"] and r["nearest_lane"]["id"] and r["inside_scene"]
    assert client.get("/api/plans/locate", headers=auth("planner"), params={"scene": "NoSuchTown", "x": 0, "y": 0}).status_code == 404
    assert client.get("/api/plans/locate", headers=auth("planner"), params={"scene": SCENE, "x": 0, "y": 0, "radius": 5000}).status_code == 422


def test_preview_gives_immediate_feedback(client, auth):
    lane = lane_xy(client, auth)["nearest_lane"]["id"]
    d = doc(rules=[{"id": "r1", "kind": "speed_limit", "x": -84.0, "y": -144.0, "params": {"speed_kmh": 25, "scope": "section"}, "target": {"lane_id": lane}}],
            scenarios=[{"id": "s1", "kind": "peak_hour", "x": -84.0, "y": -144.0, "params": {}}])
    r = client.post("/api/plans/preview", headers=auth("officer"), json=d).json()
    assert r["validation"]["ok"] and r["evidence"]["level"] == "preliminary"
    assert [e["kind"] for e in r["compiled"]["evaluated"]] == ["speed_limit"] and r["compiled"]["not_evaluated"][0]["kind"] == "peak_hour"
    assert "average travel time after" in r["pending_metrics"]
    bad = client.post("/api/plans/preview", headers=auth("officer"), json=doc(rules=[{"id": "r1", "kind": "speed_limit", "x": 9999, "y": 9999, "params": {"speed_kmh": 25}}])).json()
    assert not bad["validation"]["ok"] and bad["evidence"]["level"] == "insufficient data"
    assert client.post("/api/plans/preview", headers=auth("officer"), json=doc(scene="NoSuchTown")).status_code == 404


def test_save_versions_and_permissions(client, auth):
    d = doc(name="saved plan", props=[{"id": "p1", "kind": "warning_sign", "x": -84.0, "y": -144.0, "params": {}}])
    assert client.post("/api/plans", headers=auth("officer"), json={"plan": d}).status_code == 403
    r = client.post("/api/plans", headers=auth("planner"), json={"plan": d, "note": "first"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    d2 = {**d, "props": d["props"] + [{"id": "p2", "kind": "stop_sign", "x": -80.0, "y": -144.0, "params": {}}]}
    assert client.put(f"/api/plans/{pid}", headers=auth("operator"), json={"plan": d2}).status_code == 403  # not the owner, not planner
    r2 = client.put(f"/api/plans/{pid}", headers=auth("planner"), json={"plan": d2, "note": "added a stop sign"}, params={"base_version": 1})
    assert r2.status_code == 200 and r2.json()["latest_version"] == 2
    assert client.put(f"/api/plans/{pid}", headers=auth("planner"), json={"plan": d2}, params={"base_version": 1}).status_code == 409  # stale
    got = client.get(f"/api/plans/{pid}", headers=auth("officer")).json()
    assert got["version"] == 2 and len(got["plan"]["props"]) == 2
    old = client.get(f"/api/plans/{pid}", headers=auth("officer"), params={"version": 1}).json()
    assert len(old["plan"]["props"]) == 1 and old["note"] == "first"
    vs = client.get(f"/api/plans/{pid}/versions", headers=auth("officer")).json()
    assert [v["version"] for v in vs] == [2, 1] and vs[0]["items"]["props"] == 2
    assert pid in [p["id"] for p in client.get("/api/plans", headers=auth("officer"), params={"scene": SCENE}).json()]
    assert client.get("/api/plans/nope", headers=auth("officer")).status_code == 404
    assert client.delete(f"/api/plans/{pid}", headers=auth("planner")).status_code == 204
    assert pid not in [p["id"] for p in client.get("/api/plans", headers=auth("officer")).json()]


def wait_run(client, auth, rid, timeout=300):
    t0 = time.time()
    while time.time() - t0 < timeout:
        r = client.get(f"/api/plans/runs/{rid}", headers=auth("planner")).json()
        if r["status"] != "running":
            return r
        time.sleep(2)
    raise AssertionError("run still running")


def test_analyze_replays_on_recorded_traffic(client, auth, imported):
    lane = lane_xy(client, auth, -60.0, -125.0)["nearest_lane"]["id"]
    d = doc(name="analysis", rules=[{"id": "r1", "kind": "speed_limit", "x": -60.0, "y": -125.0, "params": {"speed_kmh": 20, "scope": "section"}, "target": {"lane_id": lane}}])
    assert client.post("/api/plans/analyze", headers=auth("officer"), json={"plan": d, "session_id": FLIGHT}).status_code == 403
    assert client.post("/api/plans/analyze", headers=auth("planner"), json={"session_id": FLIGHT}).status_code == 422
    assert client.post("/api/plans/analyze", headers=auth("planner"), json={"plan": d, "session_id": "nope"}).status_code == 404
    r = client.post("/api/plans/analyze", headers=auth("planner"), json={"plan": d, "session_id": FLIGHT})
    assert r.status_code == 202, r.text
    run = wait_run(client, auth, r.json()["id"])
    if run["status"] == "failed" and "kinematics" in (run["error"] or ""):
        return  # the fixture flight has no replayable traffic on this PC
    assert run["status"] == "done", run["error"]
    res = run["result"]
    assert res["validation"]["ok"] and res["evidence"]["level"] in ("replay-validated", "preliminary")
    if "replay" in res:  # measured on the recorded traffic, kept apart from the observed baseline
        assert res["observed"]["vehicles"] > 0 and res["replay"]["baseline"]["counted"] >= 0
        assert res["best_case"]["count"] == len(res["replay"]["removed"]) and res["worst_case"]["count"] == len(res["replay"]["added"])


def test_virtual_simulation_runs_and_returns_playable_paths(client, auth):
    lane = lane_xy(client, auth)["nearest_lane"]["id"]
    d = doc(name="virtual", rules=[{"id": "r1", "kind": "speed_limit", "x": -84.0, "y": -144.0, "params": {"speed_kmh": 15, "scope": "road"}, "target": {"lane_id": lane}}],
            scenarios=[{"id": "s1", "kind": "spawn_vehicles", "x": -84.0, "y": -144.0, "params": {"count": 3, "vehicle": "car", "speed_kmh": 40}}])
    assert client.post("/api/plans/simulate", headers=auth("officer"), json={"plan": d}).status_code == 403
    assert client.post("/api/plans/simulate", headers=auth("planner"), json={"plan": doc()}).status_code == 422  # empty
    assert client.post("/api/plans/simulate", headers=auth("planner"), json={}).status_code == 422
    bad = doc(rules=[{"id": "r1", "kind": "speed_limit", "x": 9999, "y": 9999, "params": {"speed_kmh": 25}}])
    assert client.post("/api/plans/simulate", headers=auth("planner"), json={"plan": bad}).status_code == 422
    r = client.post("/api/plans/simulate", headers=auth("planner"), json={"plan": d, "seeds": 2, "duration_s": 40})
    assert r.status_code == 202, r.text
    seen = []
    t0 = time.time()
    while time.time() - t0 < 120:  # poll fast enough to see the steps of the run
        cur = client.get(f"/api/plans/runs/{r.json()['id']}", headers=auth("planner")).json()
        if cur["status"] != "running":
            break
        if cur.get("progress"):
            seen.append(cur["progress"])
        time.sleep(0.15)
    run = wait_run(client, auth, r.json()["id"], timeout=120)
    assert run["status"] == "done", run["error"]
    assert seen and all(p["total"] == 4 and 0 <= p["done"] <= 4 and p["stage"] for p in seen)  # 2 seeds x 2 variants
    assert [p["done"] for p in seen] == sorted(p["done"] for p in seen)  # it only moves forward
    assert "progress" not in client.get(f"/api/plans/runs/{r.json()['id']}", headers=auth("planner")).json()
    res = run["result"]
    assert res["label"].startswith("virtual simulation") and res["seeds"] == 2 and res["notes"]
    m = res["metrics"]["mean_speed_kmh"]
    assert m["baseline"]["n"] == 2 and m["plan"]["n"] == 2
    tr = res["trajectories"]
    assert tr["baseline"] and tr["plan"] and len(next(iter(tr["plan"].values()))[0]) == 4


def test_report_formats_label_the_sources_and_can_be_listed_and_downloaded(client, auth):
    import base64
    import io
    lane = lane_xy(client, auth)["nearest_lane"]["id"]
    d = doc(name="report plan", rules=[{"id": "r1", "kind": "speed_limit", "x": -84.0, "y": -144.0, "params": {"speed_kmh": 25, "scope": "section"}, "target": {"lane_id": lane}}],
            scenarios=[{"id": "s1", "kind": "peak_hour", "x": -84.0, "y": -144.0, "params": {}}])
    pid = client.post("/api/plans", headers=auth("planner"), json={"plan": d}).json()["id"]
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), (80, 120, 200)).save(buf, "PNG")
    shot = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    virt = {"label": "virtual simulation (simplified model, not CARLA)", "seeds": 3, "duration_s": 60, "warmup_s": 20, "notes": ["not CARLA"], "metrics": {
        "mean_speed_kmh": {"label": "Average vehicle speed", "unit": "km/h", "baseline": {"mean": 30.0, "ci95": 1.0, "n": 3}, "plan": {"mean": 22.0, "ci95": 1.5, "n": 3}, "diff": -8.0, "verdict": "different"}}}
    body = {"plan": d, "plan_id": pid, "version": 1, "screenshots": [shot], "virtual": virt, "comparison": [
        {"name": "plan", "evidence": "replay-validated", "before": 5, "after": 8, "added": 3, "removed": 0},
        {"name": "plan / alt", "evidence": "replay-validated", "before": 5, "after": 4, "added": 0, "removed": 1}]}
    assert client.post("/api/plans/report", headers=auth("officer"), json={**body, "format": "pdf"}).status_code == 403
    assert client.post("/api/plans/report", headers=auth("planner"), json={**body, "format": "docx"}).status_code == 422
    assert client.post("/api/plans/report", headers=auth("planner"), json={**body, "format": "pdf", "screenshots": ["data:text/html;base64,AAAA"]}).status_code == 422
    out = {}
    for fmt in ("pdf", "xlsx", "csv"):
        r = client.post("/api/plans/report", headers=auth("planner"), json={**body, "format": fmt})
        assert r.status_code == 201, r.text
        rec = r.json()
        assert rec["format"] == fmt and rec["plan_id"] == pid and rec["plan_version"] == 1 and rec["stored"]
        f = client.get(rec["download_url"], headers=auth("officer"))
        assert f.status_code == 200 and len(f.content) == rec["size_bytes"]
        out[fmt] = f.content
    assert out["pdf"].startswith(b"%PDF")
    assert out["xlsx"][:2] == b"PK"
    text = out["csv"].decode("utf-8-sig")
    assert "A. Scenario overview" in text and "report plan" in text and "not simulated yet" in text
    assert "Judged by the replay?,yes" in text and "Judged by the replay?,no" in text  # the speed limit is judged, the peak-hour scenario is not
    assert "plan / alt: 4 violations after (replay)" in text  # the best compared scenario is named, with its source
    assert "replay (measured on the same recorded traffic)" in text and "rule of thumb (not measured)" in text
    assert "Recorded duration analysed,value,-" in text
    assert "Average vehicle speed (km/h),Baseline" not in text and "30 ± 1" in text and "22 ± 1.5" in text and "virtual simulation (simplified model, not CARLA)" in text  # no replay was attached, so no observed numbers are invented
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(out["xlsx"]))
    assert {"A", "B", "F"} <= set(wb.sheetnames) and "E" not in wb.sheetnames
    lst = client.get(f"/api/plans/{pid}/reports", headers=auth("officer")).json()
    assert len(lst) == 3 and {x["format"] for x in lst} == {"pdf", "xlsx", "csv"}


def test_custom_traffic_endpoint_gives_plain_traffic_and_follows_the_plan(client, auth):
    h = auth("officer")  # viewing needs no planner role
    base = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 40}).json()
    assert base["follows_plan"] is False and base["vehicles"] > 20 and base["trajectories"]
    tr = next(iter(base["trajectories"].values()))
    assert len(tr[0]) == 4 and tr[-1][0] > tr[0][0]
    lane = lane_xy(client, auth)["nearest_lane"]["id"]
    d = doc(rules=[{"id": "r1", "kind": "speed_limit", "x": -84.0, "y": -144.0, "params": {"speed_kmh": 10, "scope": "road"}, "target": {"lane_id": lane}}],
            scenarios=[{"id": "s1", "kind": "traffic_flow", "x": -84.0, "y": -144.0, "params": {"vph": 1200, "vehicle": "bus", "speed_kmh": 30}}])
    withplan = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "plan": d, "duration_s": 40}).json()
    assert withplan["follows_plan"] is True and withplan["plan_items"] == 2
    assert withplan["trajectories"] != base["trajectories"]  # the plan changed what the vehicles do
    bad = doc(rules=[{"id": "r1", "kind": "speed_limit", "x": 9999, "y": 9999, "params": {"speed_kmh": 25}}])
    assert client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "plan": bad}).status_code == 422
    assert client.post("/api/plans/traffic", headers=h, json={"scene": "NoSuchTown"}).status_code == 404
    assert client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 5}).status_code == 422


def test_traffic_schema_and_configured_traffic(client, auth):
    h = auth("officer")
    sc = client.get("/api/plans/traffic/schema", headers=h).json()
    assert {"run", "background", "mix", "drivers"} <= {s["id"] for s in sc["sections"]} and "rush_hour" in sc["presets"] and sc["defaults"]["background"]["density_per_km"] == 7.0
    assert set(sc["profiles"]) >= {"constant", "morning_peak", "event_surge"} and "freight" in sc["source_mixes"]
    light = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 40, "traffic": {"preset": "light"}}).json()
    heavy = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 40, "traffic": {"preset": "heavy"}}).json()
    assert heavy["vehicles"] > 2 * light["vehicles"] and heavy["description"].startswith("Heavy traffic") and heavy["stats"]["mean_speed_kmh"] is not None
    trucks = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 40, "traffic": {"background": {"mix": {"car": 0, "van": 0, "truck": 100, "bus": 0, "motorcycle": 0}}}}).json()
    assert set(trucks["spawned"]) == {"truck"}
    off = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "duration_s": 40, "traffic": {"preset": "empty"}}).json()
    assert off["vehicles"] == 0 and off["trajectories"] == {}
    bad = client.post("/api/plans/traffic", headers=h, json={"scene": SCENE, "traffic": {"background": {"density_per_km": 500}}})
    assert bad.status_code == 422 and "Density" in bad.text
    # the plan carries its traffic settings (saved with it) and they are validated with the plan
    d = doc(name="traffic plan", traffic={"preset": "night"})
    pid = client.post("/api/plans", headers=auth("planner"), json={"plan": d}).json()["id"]
    assert client.get(f"/api/plans/{pid}", headers=h).json()["plan"]["traffic"] == {"preset": "night"}
    pv = client.post("/api/plans/preview", headers=h, json={**d, "traffic": {"background": {"density_per_km": 500}}}).json()
    assert not pv["validation"]["ok"] and any(m["code"] == "traffic" for m in pv["validation"]["errors"])
