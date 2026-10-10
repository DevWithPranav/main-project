"""API tests (Build Plan M5) against the real PostGIS / Redis / S3 from docker-compose.yml.

Usage:
    venv\\Scripts\\python.exe -m pytest backend/tests -v -s     # -s prints the export timings
"""

import io
import json
import os
import time
import zipfile

import pytest
from websockets.sync.client import connect

SERVICE = {"Authorization": "Bearer dev-service-token"}
FLIGHT = os.getenv("TEST_FLIGHT", "20261009_201727")


# --- auth, health -------------------------------------------------------------------------------

def test_health(client):
    h = client.get("/api/health").json()
    assert {k: h[k] for k in ("db", "redis", "s3")} == {"db": "ok", "redis": "ok", "s3": "ok"}


@pytest.mark.parametrize("user,role", [("officer", "OFFICER"), ("operator", "OPERATOR"), ("planner", "PLANNER"),
                                       ("maintenance", "MAINTENANCE"), ("admin", "ADMIN")])
def test_login_and_me(client, auth, user, role):
    r = client.get("/api/me", headers=auth(user))
    assert r.json() == {"username": user, "role": role}


def test_login_rejects_bad_password(client):
    r = client.post("/api/auth/login", json={"username": "admin", "password": "nope"})
    assert r.status_code == 401 and "detail" in r.json()


def test_reads_need_a_token(client):
    assert client.get("/api/sessions").status_code == 401
    assert client.get("/api/events", headers={"Authorization": "Bearer junk"}).status_code == 401


def test_import_needs_operator(client, auth):
    for user in ("officer", "planner", "maintenance"):
        r = client.post("/api/sessions/import", headers=auth(user), json={"flight": FLIGHT})
        assert r.status_code == 403, user


# --- import, sessions, trajectories -------------------------------------------------------------

def test_import_flight(imported):
    imp = imported["import"]
    print(f"\nimport: {json.dumps({k: v for k, v in imp.items() if k != 'errors'})}")
    assert imported["session_id"] == FLIGHT and imported["town"] == "Town05" and imported["source"] == "carla"
    assert imp["n_in_file"] == imp["n_imported"] + imp["n_rejected"]
    assert imp["n_imported"] > 0 and imported["n_events"] == imp["n_imported"]
    assert imp["clips"] > 0  # the `violations` folder has render_violations.py --clips output
    assert imp["n_tracks"] > 0


def test_import_default_dir_and_reimport(client, auth, imported):
    """No violations_dir: the newest violations* folder; then back to `violations` for the rest."""
    r = client.post("/api/sessions/import", headers=auth("admin"), json={"flight": FLIGHT})
    assert r.status_code == 200, r.text
    assert "violations" in r.json()["import"]["violations_dir"]
    r = client.post("/api/sessions/import", headers=auth("operator"),
                    json={"flight": FLIGHT, "violations_dir": os.getenv("TEST_VIOLATIONS_DIR", "violations"),
                          "include_anomalies": False})
    assert r.json()["n_events"] == imported["n_events"]


def test_import_bad_flight(client, auth):
    assert client.post("/api/sessions/import", headers=auth("operator"), json={"flight": "nope"}).status_code == 404
    assert client.post("/api/sessions/import", headers=auth("admin"), json={"flight": "../x"}).status_code == 422
    assert client.post("/api/sessions/import", headers=auth("admin"),
                       json={"flight": FLIGHT, "violations_dir": "nope"}).status_code == 404


def test_sessions(client, auth, imported):
    h = auth("officer")
    assert any(s["session_id"] == FLIGHT for s in client.get("/api/sessions", headers=h).json())
    s = client.get(f"/api/sessions/{FLIGHT}", headers=h).json()
    assert set(s) >= {"session_id", "name", "source", "town", "flight", "started_at", "n_events", "profile", "scene"}
    assert client.get("/api/sessions/nope", headers=h).status_code == 404


def test_trajectories_downsampled(client, auth, imported):
    tr = client.get(f"/api/sessions/{FLIGHT}/trajectories", headers=auth("planner")).json()
    assert len(tr) == imported["import"]["n_tracks"]
    pts = max(tr.values(), key=len)
    assert all(len(p) == 4 for p in pts)
    dts = [b[0] - a[0] for a, b in zip(pts, pts[1:])]
    assert min(dts) >= 0.2 - 1e-3  # <= 5 Hz


def test_evidence_clip_served(client, auth, imported):
    items = client.get("/api/events", headers=auth("officer"), params={"session_id": FLIGHT}).json()["items"]
    url = next(e["evidence"]["clip_url"] for e in items if e["evidence"].get("clip_url"))
    r = client.get(url, headers={"Range": "bytes=0-99"})
    assert r.status_code == 206 and len(r.content) == 100 and r.headers["content-type"] in ("video/webm", "video/mp4")
    assert client.get("/api/files/nope/missing.mp4", headers=auth("officer")).status_code == 404


def test_clip_served_as_webm_once_made(client, auth):
    """Import uploads the mp4 and transcodes in the background; the WebM wins once it exists."""
    from backend.app import config, storage
    s3 = storage.client()
    s3.put_object(Bucket=config.S3_BUCKET, Key="test/clip.mp4", Body=b"mp4", ContentType="video/mp4")
    try:
        assert client.get("/api/files/test/clip.mp4", headers=auth("officer")).headers["content-type"] == "video/mp4"
        s3.put_object(Bucket=config.S3_BUCKET, Key="test/clip.webm", Body=b"webm", ContentType="video/webm")
        r = client.get("/api/files/test/clip.mp4", headers=auth("officer"))
        assert r.headers["content-type"] == "video/webm" and r.content == b"webm"
        assert set(client.get("/api/health").json()["clips"]) >= {"queued", "done", "pending"}
    finally:
        s3.delete_objects(Bucket=config.S3_BUCKET, Delete={"Objects": [{"Key": "test/clip.mp4"}, {"Key": "test/clip.webm"}]})


# --- events: filters, one, create, review -------------------------------------------------------

def test_event_filters(client, auth, imported):
    h = auth("officer")
    all_ = client.get("/api/events", headers=h, params={"session_id": FLIGHT}).json()
    assert all_["total"] == imported["n_events"] and len(all_["items"]) == all_["total"]
    for e in all_["items"]:
        assert e["session_id"] == FLIGHT and "event_id" in e
    t = all_["items"][0]["type"]
    by_type = client.get("/api/events", headers=h, params={"session_id": FLIGHT, "type": t}).json()
    assert by_type["total"] == sum(e["type"] == t for e in all_["items"])
    st = client.get("/api/events", headers=h, params={"status": "flagged,needs_review", "session_id": FLIGHT}).json()
    assert st["total"] == sum(e["status"] in ("flagged", "needs_review") for e in all_["items"])
    c = all_["items"][0]["condition"]
    assert client.get("/api/events", headers=h, params={"condition": c}).json()["total"] >= 1
    page = client.get("/api/events", headers=h, params={"session_id": FLIGHT, "limit": 3, "offset": 2}).json()
    assert len(page["items"]) == min(3, all_["total"] - 2) and page["items"][0]["event_id"] == all_["items"][2]["event_id"]
    e0 = all_["items"][0]
    box = f"{e0['x'] - 1},{e0['y'] - 1},{e0['x'] + 1},{e0['y'] + 1}"
    inbox = client.get("/api/events", headers=h, params={"bbox": box}).json()["items"]
    assert e0["event_id"] in [e["event_id"] for e in inbox]
    assert all(abs(e["x"] - e0["x"]) <= 1 and abs(e["y"] - e0["y"]) <= 1 for e in inbox)
    occ = sorted(e["occurred_at"] for e in all_["items"])
    mid = occ[len(occ) // 2]
    since = client.get("/api/events", headers=h, params={"session_id": FLIGHT, "since": mid}).json()["total"]
    until = client.get("/api/events", headers=h, params={"session_id": FLIGHT, "until": mid}).json()["total"]
    assert since == sum(o >= mid for o in occ) and until == sum(o <= mid for o in occ)
    assert client.get("/api/events", headers=h, params={"kind": "anomaly", "session_id": FLIGHT}).json()["total"] == 0
    assert client.get("/api/events", headers=h, params={"bbox": "1,2"}).status_code == 422
    one = client.get(f"/api/events/{e0['event_id']}", headers=h).json()
    assert one["event_id"] == e0["event_id"]
    assert client.get("/api/events/nope", headers=h).status_code == 404


def test_review(client, auth, imported):
    h = auth("officer")
    eid = client.get("/api/events", headers=h, params={"session_id": FLIGHT}).json()["items"][0]["event_id"]
    for user in ("operator", "planner", "maintenance"):
        assert client.post(f"/api/events/{eid}/review", headers=auth(user),
                           json={"outcome": "confirmed"}).status_code == 403
    assert client.post(f"/api/events/{eid}/review", headers=h, json={"outcome": "maybe"}).status_code == 422
    r = client.post(f"/api/events/{eid}/review", headers=h, json={"outcome": "dismissed", "note": "shadow, not a car"})
    assert r.status_code == 200
    rev = r.json()["review"]
    assert rev["outcome"] == "dismissed" and rev["by"] == "officer" and rev["note"].startswith("shadow")
    dismissed = client.get("/api/events", headers=h, params={"review": "dismissed"}).json()
    assert eid in [e["event_id"] for e in dismissed["items"]]
    none = client.get("/api/events", headers=h, params={"review": "none", "session_id": FLIGHT}).json()
    assert eid not in [e["event_id"] for e in none["items"]]
    assert none["total"] == imported["n_events"] - dismissed["total"]
    assert len(client.get(f"/api/events/{eid}/reviews", headers=h).json()) == 1
    # the served event (minus the extra occurred_at, evidence URLs are allowed) still passes the schema
    from backend.app.engine_bridge import event_errors
    ev = client.get(f"/api/events/{eid}", headers=h).json()
    ev.pop("occurred_at")
    assert event_errors([ev]) == []


def anomaly(eid: str, sid: str = "live_test") -> dict:
    return {"kind": "anomaly", "event_id": eid, "type": "pothole", "t_s": 12.5, "frame": 300, "x": 10.0, "y": -5.0,
            "severity_score": 0.7, "severity_band": "high", "confidence": 0.8, "status": "flagged", "session_id": sid}


def test_create_event_and_anomaly_review(client, auth):
    eid = f"test-anom-{time.time_ns()}"
    assert client.post("/api/events", headers=auth("officer"), json=anomaly(eid)).status_code == 403
    bad = anomaly(eid)
    bad["severity_score"] = 3
    assert client.post("/api/events", headers=SERVICE, json=bad).status_code == 422
    r = client.post("/api/events", headers=SERVICE, json=anomaly(eid))
    assert r.status_code == 201, r.text
    assert client.post("/api/events", headers=SERVICE, json=anomaly(eid)).status_code == 409
    assert client.get("/api/sessions/live_test", headers=auth("officer")).json()["source"] == "live"
    assert client.post(f"/api/events/{eid}/review", headers=auth("officer"),
                       json={"outcome": "confirmed"}).status_code == 403
    assert client.post(f"/api/events/{eid}/review", headers=auth("maintenance"),
                       json={"outcome": "confirmed"}).status_code == 200


# --- stats, exports -----------------------------------------------------------------------------

def test_stats(client, auth, imported):
    s = client.get("/api/stats", headers=auth("planner"), params={"session_id": FLIGHT}).json()
    assert set(s) >= {"by_type", "by_condition", "by_status", "by_hour", "hotspots"}
    assert sum(s["by_type"].values()) == imported["n_events"] == sum(s["by_hour"].values())
    assert s["hotspots"] and s["hotspots"][0]["count"] >= s["hotspots"][-1]["count"]


@pytest.mark.parametrize("fmt", ["csv", "geojson", "xlsx", "pdf"])
def test_export(client, auth, imported, fmt):
    t0 = time.perf_counter()
    r = client.get("/api/export", headers=auth("officer"), params={"format": fmt, "session_id": FLIGHT})
    dt = time.perf_counter() - t0
    print(f"\nexport {fmt}: {dt:.3f} s round trip, server {r.headers.get('x-export-seconds')} s, "
          f"{len(r.content)} bytes, {r.headers.get('x-event-count')} events")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert int(r.headers["x-event-count"]) == imported["n_events"]
    assert dt <= 10.0
    if fmt == "csv":
        lines = r.content.decode("utf-8-sig").strip().splitlines()
        assert len(lines) == imported["n_events"] + 1 and lines[0].startswith("event_id,")
    elif fmt == "geojson":
        fc = json.loads(r.content)
        assert fc["type"] == "FeatureCollection" and len(fc["features"]) == imported["n_events"]
    elif fmt == "xlsx":
        assert zipfile.is_zipfile(io.BytesIO(r.content))
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert wb["Events"].max_row == imported["n_events"] + 1
    else:
        assert r.content.startswith(b"%PDF")


def test_export_bad_format(client, auth):
    assert client.get("/api/export", headers=auth("officer"), params={"format": "doc"}).status_code == 422


# --- profiles, conditions, scenes ---------------------------------------------------------------

def test_profiles_seeded(client, auth):
    names = [p["name"] for p in client.get("/api/profiles", headers=auth("officer")).json()]
    assert "default" in names
    assert client.get("/api/profiles/default", headers=auth("officer")).json()["name"] == "default"
    assert client.get("/api/profiles/nope", headers=auth("officer")).status_code == 404


def test_profile_put_validation_and_history(client, auth):
    doc = client.get("/api/profiles/default", headers=auth("admin")).json()
    assert client.put("/api/profiles/default", headers=auth("officer"),
                      json={"profile": doc, "note": "x"}).status_code == 403
    bad = json.loads(json.dumps(doc))
    bad["violations"]["speeding"]["params"]["no_such_threshold"] = 1
    r = client.put("/api/profiles/default", headers=auth("planner"), json={"profile": bad, "note": "bad"})
    assert r.status_code == 422 and "no_such_threshold" in r.json()["detail"]
    bad2 = {k: v for k, v in doc.items() if k != "profile_version"}
    assert client.put("/api/profiles/default", headers=auth("planner"), json={"profile": bad2}).status_code == 422
    assert client.put("/api/profiles/default", headers=auth("planner"),
                      json={"profile": {**doc, "name": "other"}}).status_code == 422
    before = client.get("/api/profiles/default/history", headers=auth("officer")).json()
    good = json.loads(json.dumps(doc))
    good["violations"]["speeding"]["enabled"] = False
    r = client.put("/api/profiles/default", headers=auth("planner"), json={"profile": good, "note": "test: speeding off"})
    assert r.status_code == 200, r.text
    assert r.json()["version"] == before[0]["version"] + 1 and r.json()["by"] == "planner"
    hist = client.get("/api/profiles/default/history", headers=auth("officer")).json()
    assert len(hist) == len(before) + 1 and hist[0]["note"] == "test: speeding off" and hist[0]["by"] == "planner"
    prof = client.get("/api/profiles/default", headers=auth("officer")).json()
    assert prof["violations"]["speeding"]["enabled"] is False


def test_conditions_and_scenes(client, auth):
    h = auth("officer")
    assert len(client.get("/api/conditions", headers=h).json()["conditions"]) == 34
    scenes = client.get("/api/scenes", headers=h).json()
    assert "Town05" in [s["town"] for s in scenes]
    assert client.get("/api/scenes/town05", headers=h).json()["lanes"]
    assert client.get("/api/scenes/nope", headers=h).status_code == 404


# --- planning -----------------------------------------------------------------------------------

def test_recommendations_and_planner_history(client, auth):
    recs = client.get("/api/recommendations", headers=auth("planner"))
    assert recs.status_code == 200 and isinstance(recs.json(), list)
    rid = recs.json()[0]["id"] if recs.json() else "rec-test"
    assert client.post(f"/api/recommendations/{rid}/decision", headers=auth("officer"),
                       json={"decision": "accepted", "rationale": "x"}).status_code == 403
    assert client.post(f"/api/recommendations/{rid}/decision", headers=auth("planner"),
                       json={"decision": "maybe"}).status_code == 422
    r = client.post(f"/api/recommendations/{rid}/decision", headers=auth("planner"),
                    json={"decision": "rejected", "rationale": "budget"})
    assert r.status_code == 200
    hist = client.get("/api/planner/history", headers=auth("officer")).json()
    assert hist[0]["recommendation_id"] == rid and hist[0]["by"] == "planner" and hist[0]["rationale"] == "budget"


# --- live ---------------------------------------------------------------------------------------

def vehicle(sid: str, tid: int, t: float) -> dict:
    return {"session_id": sid, "track_id": tid, "cls": "car", "x": 1.0 + t, "y": 2.0, "speed_kmh": 30.0,
            "heading_deg": 90.0, "lane_id": "r1_s0_l1", "state": "ok", "t_s": t}


def recv_until(ws, pred, timeout=5.0):
    end = time.time() + timeout
    while time.time() < end:
        try:
            m = json.loads(ws.recv(timeout=max(0.05, end - time.time())))
        except TimeoutError:
            break
        if pred(m):
            return m
    raise AssertionError("no matching message")


def ws_url(server: str, query: str) -> str:
    return f"{server.replace('http', 'ws', 1)}/api/ws/live?{query}"


def test_live_state_needs_service_token(client, auth):
    assert client.post("/api/live/state", headers=auth("operator"), json=[vehicle("s", 1, 0)]).status_code == 403
    assert client.post("/api/live/state", headers=SERVICE, json=[{"session_id": "s"}]).status_code == 422
    bad = vehicle("s", 1, 0)
    bad["state"] = "weird"
    assert client.post("/api/live/state", headers=SERVICE, json=[bad]).status_code == 422


def test_ws_rejects_bad_token(server):
    with pytest.raises(Exception):
        with connect(ws_url(server, "token=junk")) as ws:
            ws.recv(timeout=2)


def test_live_state_and_ws(client, auth, server):
    sid = f"ws_{time.time_ns()}"
    tok = auth("officer")["Authorization"][7:]
    with connect(ws_url(server, f"token={tok}&session_id={sid}")) as ws:
        time.sleep(0.3)  # let the server register the socket
        r = client.post("/api/live/state", headers=SERVICE, json=[vehicle(sid, 1, 0.0), vehicle(sid, 2, 0.0)])
        assert r.json() == {"received": 2}
        m = recv_until(ws, lambda m: m["type"] == "vehicles")
        assert {v["track_id"] for v in m["items"]} == {1, 2} and m["items"][0]["session_id"] == sid
        # another session's vehicles are not sent to this socket
        client.post("/api/live/state", headers=SERVICE, json=[vehicle(sid + "_other", 9, 0.0)])
        # a burst of 30 batches reaches the socket at <= 10 Hz, the newest batch last
        t0 = time.time()
        for i in range(30):
            client.post("/api/live/state", headers=SERVICE, json=[vehicle(sid, 1, 1.0 + i)])
        got = []
        while True:
            try:
                got.append(json.loads(ws.recv(timeout=0.6)))
            except TimeoutError:
                break
        veh = [m for m in got if m["type"] == "vehicles"]
        elapsed = time.time() - t0
        print(f"\nws: 30 posts in {elapsed:.2f} s -> {len(veh)} vehicle messages")
        assert veh and veh[-1]["items"][0]["t_s"] == 30.0
        assert all(v["items"][0]["session_id"] == sid for v in veh)
        assert len(veh) <= elapsed * 10 + 2
        # new events are pushed too
        eid = f"ws-anom-{time.time_ns()}"
        assert client.post("/api/events", headers=SERVICE, json=anomaly(eid, sid)).status_code == 201
        m = recv_until(ws, lambda m: m["type"] == "event")
        assert m["event"]["event_id"] == eid
    # vehicle state stays in Redis for 2 s: a new socket gets the snapshot first
    with connect(ws_url(server, f"token={tok}&session_id={sid}")) as ws2:
        m = recv_until(ws2, lambda m: m["type"] == "vehicles", timeout=1.5)
        assert 1 in {v["track_id"] for v in m["items"]}
